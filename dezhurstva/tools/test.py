#!/usr/bin/env python3
"""Проверка приложения в Chromium: вкладки, выбор себя, сетка, пожелания, ширины экрана, ошибки JS.
Запуск из корня репозитория: python3 dezhurstva/tools/test.py [папка для скриншотов]"""
import functools, http.server, json, pathlib, sys, threading
from playwright.sync_api import sync_playwright

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "tools" / "shots"
OUT.mkdir(parents=True, exist_ok=True)
class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass


srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(Quiet, directory=str(ROOT)))
threading.Thread(target=srv.serve_forever, daemon=True).start()
URL = f"http://127.0.0.1:{srv.server_address[1]}/index.html"
data = json.loads((ROOT / "data" / "schedule.json").read_text())
fails, errors = [], []


def text(loc):
    return loc.inner_text().replace("\xa0", " ")


def check(cond, msg):
    print(("ok  " if cond else "FAIL"), msg)
    if not cond:
        fails.append(msg)


# Меню «Поделиться» вместо настоящего: запоминаем, какие файлы в него ушли
SHARE_STUB = """window.__shared = []; navigator.canShare = d => !!(d && d.files);
navigator.share = async d => { for (const f of d.files) { const u = new Uint8Array(await f.arrayBuffer());
  window.__shared.push([f.name, f.type, String.fromCharCode(...u.slice(0, 5)), u.length]); } };"""


def page_at(b, when, width=390, height=844, me=None, tab=None, init=""):
    ctx = b.new_context(viewport={"width": width, "height": height}, timezone_id="Asia/Yekaterinburg", locale="ru-RU", device_scale_factor=2)
    pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    pg.on("console", lambda m: errors.append(m.text) if m.type == "error" else None)
    pg.clock.install(time=when)
    if me:
        init += f"localStorage.setItem('dezh:me', {json.dumps(me)});"
    if tab:
        init += f"localStorage.setItem('dezh:tab', {json.dumps(tab)});"
    if init:
        pg.add_init_script(init)
    pg.goto(URL)
    pg.wait_for_selector("#view > *")
    return ctx, pg


def pdf_read(body):
    """pdfinfo и текст PDF (poppler)."""
    import subprocess, tempfile
    with tempfile.NamedTemporaryFile(suffix=".pdf") as f:
        f.write(body); f.flush()
        return (subprocess.run(["pdfinfo", f.name], capture_output=True, text=True).stdout,
                subprocess.run(["pdftotext", "-layout", f.name, "-"], capture_output=True, text=True).stdout)


def hscroll(pg):
    return pg.evaluate("""() => { const m = document.querySelector('main'); const bad = [];
      for (const el of document.querySelectorAll('#view *')) { if (el.closest('.chips')) continue; const r = el.getBoundingClientRect();
        if (r.width && (r.right > innerWidth + 1 || r.left < -1)) bad.push(el.tagName + '.' + el.className + ' ' + Math.round(r.right)); }
      return { doc: document.documentElement.scrollWidth > innerWidth, main: m.scrollWidth > m.clientWidth, bad: bad.slice(0, 5) }; }""")


def split_words(pg):
    """Слова, разорванные между строками (перенос внутри слова)."""
    return pg.evaluate("""() => { const bad = [], range = document.createRange();
      const w = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT); let node;
      while ((node = w.nextNode())) { const el = node.parentElement;
        if (!el || el.closest('script, style, textarea, [hidden], .chips') || !el.getClientRects().length) continue;
        const re = /[A-Za-zА-Яа-яЁё0-9]{2,}/g; let m;
        while ((m = re.exec(node.textContent))) { range.setStart(node, m.index); range.setEnd(node, m.index + m[0].length);
          const tops = new Set([...range.getClientRects()].filter(r => r.width > 0).map(r => Math.round(r.bottom)));
          if (tops.size > 1) bad.push(m[0]); } }
      return bad.slice(0, 8); }""")


def clipped(pg):
    return pg.evaluate("""() => { const out = [];
      const sel = '.tab, .brush button, .seg button, .wdrow button, .btn, .chip, .mbar h2, .day .w, .day .n, .wcal .day .t, .pill, .stepper output, .dlist b.nm';
      for (const el of document.querySelectorAll(sel)) { const st = getComputedStyle(el);
        if (st.display === 'none' || !el.getClientRects().length) continue;
        if (el.scrollWidth > el.clientWidth + 1) out.push((el.className || el.tagName) + ': ' + el.textContent.trim().slice(0, 20));
        const pr = el.parentElement.getBoundingClientRect(), r = el.getBoundingClientRect();
        if (el.matches('.day .w, .day .n, .wcal .day .t, .dlist b.nm') && (r.left < pr.left - 1 || r.right > pr.right + 1)) out.push('клетка: ' + el.textContent.trim()); }
      return out.slice(0, 6); }""")


with sync_playwright() as p:
    b = p.chromium.launch()
    # 1. Первый запуск: выбор себя
    ctx, pg = page_at(b, "2026-11-05T12:00:00+05:00")
    check(pg.locator(".name").count() == len(data["doctors"]), "выбор себя: все врачи")
    pg.screenshot(path=str(OUT / "01-pick.png"))
    pg.locator(".name", has_text="Клепиков").click()
    hero = text(pg.locator(".hero"))
    check("сегодня в 16:18" in hero and "5 ноября" in hero, f"ближайшее дежурство Клепикова: {hero!r}")
    check(pg.locator(".dlist li:not(.mark)").count() == 7, "7 дежурств в ноябре")
    check("134 ч 48 мин" in text(pg.locator(".sum")), "сумма часов 134 ч 48 мин")
    check(pg.locator('a[href^="webcal:"]').get_attribute("href").endswith("/cal/klepikov.ics"), "ссылка webcal")
    pg.screenshot(path=str(OUT / "02-me.png"), full_page=False)
    # 2. График: сейчас дежурит (5 ноября 03:00 — Алекешова с 4 ноября, праздник, до 08:00)
    ctx.close()
    ctx, pg = page_at(b, "2026-11-05T03:00:00+05:00", me="klepikov", tab="all")
    now = pg.locator(".now").inner_text()
    check("Сейчас" in now and "Алекешова" in now and "до 08:00" in now, f"сейчас дежурит: {now!r}")
    check(pg.locator(".cal .day").count() == 30, "30 дней в сетке")
    check(pg.locator(".cal .day.off").count() == 10, "выходных и праздников 10 (9 сб/вс + 4 ноября)")
    check(pg.locator(".cal .day.hol").count() == 1, "праздник 4 ноября отмечен")
    check(pg.locator(".cal .day.mine").count() == 7, "мои дни подсвечены")
    codes = pg.eval_on_selector_all(".cal .day .w", "els => [...new Set(els.map(e => e.textContent))]")
    check(all(len(c) == 2 and c[0].isupper() and c[1].islower() for c in codes) and len(codes) == len(set(codes)) == 7
          and {"Лк", "Ло", "Кл", "Ка", "Ом"} <= set(codes), f"в сетке по две буквы, без совпадений: {sorted(codes)}")
    gap = pg.evaluate("""() => { let worst = 99; for (const d of document.querySelectorAll('.cal .day')) {
        const n = d.querySelector('.n'), w = d.querySelector('.w'); if (!w) continue;
        worst = Math.min(worst, w.getBoundingClientRect().top - n.getBoundingClientRect().bottom); } return worst; }""")
    check(gap >= 0, f"число и буквы в клетке не наезжают (зазор {gap:.1f}px)")
    pg.screenshot(path=str(OUT / "03-grid.png"), full_page=False)
    pg.locator('.cal .day[data-k="2026-11-01"]').click()
    dc = pg.locator(".daycard").inner_text()
    check("Лекерова" in dc and "08:00 → 08:00" in dc and "сутки" in dc, f"карточка дня 1 ноября: {dc!r}")
    pg.locator('.chip[data-id="leonard"]').click()
    check(pg.locator(".cal .day.hit").count() == 4, "подсветка Леонарда: 4 пятницы")
    pg.locator('.seg button[data-v="list"]').click()
    check(pg.locator(".dlist li").count() == 4, "список с фильтром: 4 строки")
    pg.locator('.chip[data-id=""]').click()
    one = pg.evaluate("""() => [...document.querySelectorAll('.dlist b.nm')].every(b => b.getClientRects().length === 1 && b.scrollWidth <= b.clientWidth + 1)""")
    check(one, "в списке фамилии в одну строку, без переноса")
    pg.screenshot(path=str(OUT / "04-list.png"), full_page=False)
    pg.locator('.seg button[data-v="load"]').click()
    nt = text(pg.locator(".norm"))
    check("Часов в неделю\n39" in nt and "Рабочих дней\n20" in nt and "Сокращено, ч\n1" in nt and "155" in nt and "3 ноября" in nt,
          f"норма ноября: 39 ч, 20 дней, −1 ч, 155 ч {nt[:120]!r}")
    shares = pg.eval_on_selector_all(".load li", "els => els.map(e => e.querySelector('.nm').textContent + ' ' + e.querySelector('.sh').textContent)")
    check(len(shares) == len(data["doctors"]) and "Клепиков 0,87" in shares, f"доли ставки: {shares}")
    check("3,57 ставки" in text(pg.locator(".load ~ .sum")), "итого 3,57 ставки")
    pg.screenshot(path=str(OUT / "04b-load.png"), full_page=False)
    ctx.close()
    # 3. Пожелания (сегодня 12 октября → декабрь, до 15 ноября)
    ctx, pg = page_at(b, "2026-10-12T21:00:00+05:00", me="omutov", tab="wish")
    h = text(pg.locator("#view .card").first)
    check("декабрь 2026" in h and "15 ноября" in h, f"пожелания на декабрь до 15 ноября: {h[:90]!r}")
    check(pg.locator('.wcal .day[data-k="2026-12-31"]').get_attribute("class").count("off") == 1, "31 декабря 2026 — выходной по переносу")
    pg.locator('.brush button[data-v="no"]').click()
    pg.locator('.toggle[data-act="range"]').click()
    pg.locator('.wcal .day[data-k="2026-12-01"]').click()
    pg.locator('.wcal .day[data-k="2026-12-05"]').click()
    pg.locator('.toggle[data-act="range"]').click()
    pg.locator('.brush button[data-v="yes"]').click()
    pg.locator('.wcal .day[data-k="2026-12-19"]').click()
    pg.locator('.wdrow.no button[data-i="0"]').click()
    pg.locator('.wdrow.yes button[data-i="4"]').click()
    pg.locator('button[data-act="qmode"][data-f="total"][data-v="max"]').click()
    pg.locator('button[data-act="qn"][data-f="total"][data-d="1"]').click()
    pg.fill("#note", "Не ставить подряд с выходными")
    s = pg.locator("#summary").inner_text()
    check("Не ставить: пн; 1–5" in s and "Хочу: пт; 19" in s and "всего не больше 5" in s and "Комментарий: Не ставить подряд" in s, f"итог пожеланий: {s!r}")
    cfg = json.loads((ROOT / "data" / "config.json").read_text())["form"]
    if cfg.get("action"):
        posted = []
        pg.route(cfg["action"], lambda route: (posted.append(route.request.post_data or ""), route.fulfill(status=200, body="ok")))
        pg.locator('button[data-act="send"]').click()
        pg.wait_for_selector(".sent")
        from urllib.parse import parse_qs
        body = parse_qs(posted[0]) if posted else {}
        val = (body.get(cfg["entry"]) or [""])[0]
        check(len(posted) == 1 and "#DEZH " in val and '"who":"omutov"' in val and '"m":"2026-12"' in val,
              f"отправка в Google Форму: поле {cfg['entry']}, {len(val)} симв.")
    else:
        check(pg.locator('button[data-act="send"]').is_disabled(), "отправка выключена, пока форма не настроена")
    pg.screenshot(path=str(OUT / "05-wish.png"), full_page=True)
    ctx.close()
    # 4. Excel
    ctx, pg = page_at(b, "2026-10-12T21:00:00+05:00", me="omutov", tab="files", init=SHARE_STUB)
    btns = pg.eval_on_selector_all('.file button[data-act="fshare"]', "els => els.map(e => [e.dataset.kind, e.textContent])")
    check(btns == [["pdf", "PDF — печать"], ["xlsx", "Excel"]] and pg.locator(".file a").count() == 0,
          f"у месяца кнопки PDF и Excel, не ссылки (ссылка в приложении с экрана «Домой» открывала файл без пути назад): {btns}")
    hrefs = [pg.evaluate(f"pubFile('2026-11', '{k}')[0]") for k in ("pdf", "xlsx")]
    check(hrefs[0].startswith("files/2026-11.pdf?v=") and hrefs[1].startswith("files/2026-11.xlsx"), f"адреса PDF и Excel ноября: {hrefs}")
    r = pg.request.get(URL.replace("index.html", hrefs[1]))
    check(r.ok and len(r.body()) > 3000, "файл Excel скачивается")
    r = pg.request.get(URL.replace("index.html", hrefs[0]))
    check(r.ok and r.body()[:5] == b"%PDF-", "PDF скачивается")
    info, txt = pdf_read(r.body())
    check("Pages:           1" in info and "841.92 x 594.96" in info, "PDF — один лист A4 альбомом")
    check("ГРАФИК ДЕЖУРСТВ ПСО — НОЯБРЬ 2026" in txt and "КЛЕПИКОВ" in txt and "телефон" not in txt, "в PDF заголовок, фамилии, без телефонов")
    pg.wait_for_function("Object.keys(got).length === 2")
    url0 = pg.url
    pg.click('.file button[data-kind="pdf"]')
    pg.wait_for_function("__shared.length === 1")
    pg.click('.file button[data-kind="xlsx"]')
    pg.wait_for_function("__shared.length === 2")
    sh = pg.evaluate("__shared")
    check(sh[0][:3] == ["График дежурств ПСО — ноябрь 2026.pdf", "application/pdf", "%PDF-"] and sh[1][0] == "График дежурств ПСО — ноябрь 2026.xlsx"
          and sh[1][2][:2] == "PK" and pg.url == url0, f"файлы скачаны заранее и уходят в меню «Поделиться», страница остаётся: {[s[:2] for s in sh]}")
    pg.screenshot(path=str(OUT / "06-files.png"), full_page=False)
    ctx.close()
    # Врачи, исключённые в config.json, нигде не показываются
    excl = json.loads((ROOT / "data" / "config.json").read_text()).get("exclude", [])
    ctx, pg = page_at(b, "2026-11-05T12:00:00+05:00", tab="me")
    body = pg.inner_text("body")
    check(not any(x in body for x in ("Давыдова", "ДАВЫДОВА")) and "davydova" not in data["doctors"] and "davydova" in excl, "Давыдовой нет в приложении")
    ctx.close()
    # Размер текста iPhone (Dynamic Type): -apple-system-body 23 px вместо 17 → шрифт приложения ×1,35
    ctx = b.new_context(viewport={"width": 390, "height": 844})
    pg = ctx.new_page()
    pg.add_init_script("""(() => { const s = CSS.supports.bind(CSS);
      CSS.supports = (a, v) => (a === 'font' && v === '-apple-system-body') ? true : s(a, v);
      const g = window.getComputedStyle;
      window.getComputedStyle = function (el, ...r) { const cs = g.call(window, el, ...r);
        if (el && el.textContent === 'x' && el.style && el.style.visibility === 'hidden')
          return new Proxy(cs, { get: (o, k) => k === 'fontSize' ? '23px' : (typeof o[k] === 'function' ? o[k].bind(o) : o[k]) });
        return cs; }; })()""")
    pg.goto(URL); pg.wait_for_selector("#view > *")
    fs = pg.evaluate("[document.documentElement.style.fontSize, getComputedStyle(document.documentElement).fontSize]")
    check(fs[0] == "142.06%", f"крупный текст на iPhone увеличивает шрифт приложения: {fs}")
    ctx.close()
    ctx, pg = page_at(b, "2026-11-05T12:00:00+05:00", tab="me")
    fs = pg.evaluate("[document.documentElement.style.fontSize, getComputedStyle(document.documentElement).fontSize]")
    check(fs[0] in ("105%", "105.00%") and fs[1] == "16.8px", f"обычный размер текста — шрифт как раньше: {fs}")
    ctx.close()
    # 4а. Для печати: телефоны из файла графика (вымышленные), сборка файла в браузере
    import base64, io, subprocess, tempfile, zipfile
    tmp = pathlib.Path(tempfile.mkdtemp())
    fake = {d: f"8900000{i:04d}" for i, d in enumerate(data["doctors"])}
    (tmp / "ph.json").write_text(json.dumps(fake))
    subprocess.run([sys.executable, str(ROOT / "tools" / "make_xlsx.py"), "2026-11", "--full", str(tmp / "src.xlsx"),
                    "--phones", str(tmp / "ph.json")], check=True, capture_output=True)
    M11o = data["months"]["2026-11"]["order"]
    ctx, pg = page_at(b, "2026-10-12T21:00:00+05:00", me="klepikov", tab="files", init=SHARE_STUB)
    pg.locator("details.org > summary").click()
    pg.set_input_files("#phfile", str(tmp / "src.xlsx"))
    pg.wait_for_function("document.querySelectorAll('[data-ph]')[0].value !== ''")
    vals = pg.eval_on_selector_all("[data-ph]", "els => els.map(e => [e.dataset.ph, e.value])")
    check(all(v == fake[k] for k, v in vals) and len(vals) == len(fake), "телефоны взяты из файла графика")
    # Файл «как попало»: номер в соседней ячейке, +7 (…), без восьмёрки, числом, с припиской, ячейка из частей с разным шрифтом
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    wbx = openpyxl.Workbook(); wsx = wbx.active
    wsx.append(["График", None, None])
    wsx.append(["1. Клепиков Е.В.", "+7 (900) 000-00-01", "+"])
    wsx.append([CellRichText([TextBlock(InlineFont(b=True), "КАРДАШИНА "), "89000000002"]), None, "+"])
    wsx.append(["Лекерова", 89000000003, None])
    wsx.append(["Хаматнурова, моб.", "900 000 00 04", None])
    wsx.append(["ЛЕОНАРД 8-900-000-00-05 (моб.)", None, "+"])
    wsx.append(["Алекешова", "тел.: 8 (900) 000 00 06", None])
    wsx.append(["Лекерова вместо Омутова 01.11", None, None])
    wsx.append(["Омутов", "8 900 000 00 07", None])
    wbx.save(tmp / "messy.xlsx")
    pg.set_input_files("#phfile", str(tmp / "messy.xlsx"))
    pg.wait_for_function("document.querySelector('[data-ph=\"omutov\"]').value.includes('07')")
    got = dict(pg.eval_on_selector_all("[data-ph]", "els => els.map(e => [e.dataset.ph, e.value.replace(/\\D/g, '')])"))
    want = {"klepikov": "79000000001", "kardashina": "89000000002", "lekerova": "89000000003", "khamatnurova": "89000000004",
            "leonard": "89000000005", "alekeshova": "89000000006", "omutov": "89000000007"}
    check(got == want, f"телефоны из «неудобного» файла: {got}")
    pg.set_input_files("#phfile", str(tmp / "src.xlsx"))
    pg.wait_for_function("document.querySelector('[data-ph=\"omutov\"]').value === '" + fake["omutov"] + "'")
    pg.fill('[data-ph="omutov"]', "8 900 111-22-33")
    pg.locator('button[data-act="pprep"]').first.click()
    pg.wait_for_selector('button[data-act="pshare"]')
    check(pg.eval_on_selector_all('button[data-act="pshare"]', "els => els.map(e => e.textContent)") == ["PDF — печать", "Excel"],
          "готовы PDF и Excel для печати")
    got = pg.evaluate("""async () => Promise.all(printReady.out.map(async o => { const buf = new Uint8Array(await o.file.arrayBuffer()); let s = '';
      for (let i = 0; i < buf.length; i++) s += String.fromCharCode(buf[i]); return [o.kind, o.file.name, btoa(s)]; }))""")
    pg.click('button[data-act="pshare"][data-k="0"]')
    pg.wait_for_function("__shared.length === 1")
    sh = pg.evaluate("__shared")[0]
    check(sh[:3] == ["График дежурств ПСО — ноябрь 2026 (для печати).pdf", "application/pdf", "%PDF-"], f"PDF с телефонами уходит в меню «Поделиться»: {sh[:2]}")
    pdf_b = base64.b64decode(got[0][2])
    b64 = got[1][2]
    check(got[0][:2] == ["PDF", "График дежурств ПСО — ноябрь 2026 (для печати).pdf"] and got[1][1].endswith("(для печати).xlsx"), "имена файлов для печати")
    info, txt = pdf_read(pdf_b)
    check("Pages:           1" in info and "Врач, телефон" in txt and fake["klepikov"] in txt and "8 900 111-22-33" in txt
          and all(fake[d] in txt for d in M11o if d != "omutov"), "PDF для печати: один лист, телефоны у всех фамилий")
    (tmp / "print.pdf").write_bytes(pdf_b)
    subprocess.run(["pdftoppm", "-r", "50", "-png", "-singlefile", str(tmp / "print.pdf"), str(OUT / "07-print-pdf")], check=True)
    z = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64)))
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64))).worksheets[0]
    a_col = [ws.cell(r, 1).value for r in range(3, 13)]
    check(a_col[0] == "Врач, телефон" and a_col[2] == f"КЛЕПИКОВ {fake['klepikov']}" and "ОМУТОВ 8 900 111-22-33" in a_col,
          f"файл для печати: телефоны у фамилий {a_col[:3]}")
    check(ws.cell(5, 6).value == "+", "в файле для печати те же «+»")
    wv = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)), data_only=True).worksheets[0]
    M11 = data["months"]["2026-11"]
    want = [sum(d in v for v in M11["duty"].values()) for d in M11["order"]]
    last = 4 + len(want) + 1
    check([wv.cell(r, 32).value for r in range(5, 5 + len(want))] == want and all(wv.cell(last, c).value == 1 for c in range(2, 32))
          and wv.cell(last, 32).value == 30, f"в файле для печати посчитаны «Всего» {want} и «Дежурных за день» (видно и без пересчёта)")
    check(wv.max_row == last, f"в файле для печати под таблицей ничего нет (строк {wv.max_row})")
    pg.screenshot(path=str(OUT / "07-print.png"), full_page=False)
    ctx.close()
    for f in tmp.iterdir():
        f.unlink()
    tmp.rmdir()
    # 6. Согласование: подставляем предварительный декабрь (версия 2; организатор уже согласен)
    import re as _re
    html = (ROOT / "index.html").read_text()
    mm = _re.search(r'(<script id="data" type="application/json">)(.*?)(</script>)', html, _re.S)
    DD = json.loads(mm.group(2).replace("<\\/", "</"))
    order = DD["months"]["2026-11"]["order"]
    dec_days = [f"2026-12-{d:02d}" for d in range(1, 32)]
    DD["months"]["2026-12"] = {"title": "Декабрь 2026", "order": order, "shift": DD["months"]["2026-11"]["shift"],
        "off": DD["cal"]["2026-12"]["off"], "duty": {k: [order[i % len(order)]] for i, k in enumerate(dec_days)}, "marks": [],
        "status": "preliminary", "version": 2,
        "confirm": {"version": 2, "published": "2026-11-20T05:00:00Z", "need": order, "confirmed": ["klepikov"], "disagreed": 0}}
    DD["files"]["2026-12"] = DD["files"]["2026-11"]
    pre_html = html.replace(mm.group(2), json.dumps(DD, ensure_ascii=False).replace("</", "<\\/"))

    def pre_page(me, width=390):
        ctx = b.new_context(viewport={"width": width, "height": 844}, timezone_id="Asia/Yekaterinburg", locale="ru-RU")
        ctx.route("**/index.html", lambda r: r.fulfill(body=pre_html, content_type="text/html; charset=utf-8"))
        pg = ctx.new_page()
        pg.on("pageerror", lambda e: errors.append(str(e)))
        pg.clock.install(time="2026-11-21T10:00:00+05:00")
        pg.add_init_script(f"localStorage.setItem('dezh:me', '{me}'); localStorage.setItem('dezh:tab', 'all')")
        pg.goto(URL)
        pg.wait_for_selector("#view > *")
        return ctx, pg

    ctx, pg = pre_page("omutov")
    check(pg.is_visible("#ackSheet") and "декабрь 2026 (версия 2)" in pg.inner_text("#ackTitle"), "при открытии — уведомление о предварительном графике")
    check(pg.locator('.tab[data-tab="me"] .badge').count() == 1, "точка на вкладке «Мои»")
    pg.click("#ackOpen")
    card = text(pg.locator(".ack"))
    check("Согласны 1 из 7" in card and pg.locator(".ack .mine span").count() >= 4, f"карточка: свои дежурства и «согласны 1 из 7» {card[:80]!r}")
    cfg = json.loads((ROOT / "data" / "config.json").read_text())["form"]
    posted = []
    pg.route(cfg["action"], lambda route: (posted.append(route.request.post_data or ""), route.fulfill(status=200, body="ok")))
    pg.click('.ack [data-act="ackno"]')
    pg.click('.ack [data-act="acksend"]')
    check(not posted and pg.is_visible("[data-acknote]"), "без причины «Не согласен» не отправляется")
    pg.fill("[data-acknote]", "Не могу 24-го — конференция")
    pg.click('.ack [data-act="acksend"]')
    pg.wait_for_selector(".ack .done.no")
    from urllib.parse import parse_qs
    val = (parse_qs(posted[-1]).get(cfg["entry"]) or [""])[0] if posted else ""
    check('"kind":"ack"' in val and '"ok":false' in val and '"ver":2' in val and "конференция" in val and "НЕ СОГЛАСЕН" in val,
          "несогласие с причиной ушло организатору (в форму)")
    check(pg.locator('.tab[data-tab="me"] .badge').count() == 0, "после ответа точка пропала")
    shot = OUT / "08-ack.png"
    pg.click('.ack [data-act="ackchange"]')
    pg.click('.ack [data-act="acksend"][data-ok="1"]')
    pg.wait_for_selector(".ack .done:not(.no)")
    val = (parse_qs(posted[-1]).get(cfg["entry"]) or [""])[0]
    check('"ok":true' in val and "СОГЛАСЕН" in val, "«Изменить ответ» → «Согласен» отправлено")
    pg.evaluate("document.querySelector('main').scrollTop = 0")
    pg.screenshot(path=str(shot))
    for w in (320, 390):
        pg.set_viewport_size({"width": w, "height": 844})
        r, c, sw = hscroll(pg), clipped(pg), split_words(pg)
        check(not r["bad"] and not c and not sw, f"{w}px карточка согласования помещается {r['bad'] or ''} {c or ''} {sw or ''}")
    pg.click('.tab[data-tab="all"]')
    pg.locator('.mbar [data-act="next"]').click()
    st = text(pg.locator(".status").first)
    check("предварительный · версия 2" in st and "согласны 1 из 7" in st, f"в «Графике» статус месяца: {st!r}")
    pg.click('.tab[data-tab="files"]')
    check(pg.locator(".file .pill.pre").count() == 1, "в «Excel» у декабря пометка «предварительный»")
    ctx.close()
    ctx, pg = pre_page("klepikov")
    check(not pg.is_visible("#ackSheet") and pg.locator('.tab[data-tab="me"] .badge').count() == 0, "у организатора (уже согласен) уведомления нет")
    pg.click('.tab[data-tab="me"]')
    check("Ваше согласие учтено" in text(pg.locator(".ack")), "организатору — «согласие учтено»")
    ctx.close()
    # 5. Ширины: ничего не вылезает за экран и не обрезается внутри кнопок и клеток; ещё раз с текстом крупнее на 15%
    for scale in ("", "120.75%", "147%"):
        for w in (320, 360, 375, 390, 393, 402, 430, 440, 820, 1180):
            for tab in ("me", "all", "wish", "files"):
                ctx, pg = page_at(b, "2026-11-05T12:00:00+05:00" if tab != "wish" else "2026-10-12T21:00:00+05:00", width=w, height=900, me="khamatnurova", tab=tab)
                if scale:
                    pg.evaluate(f"document.documentElement.style.fontSize = '{scale}'")
                if tab == "files":
                    pg.locator("details.org > summary").click()
                if tab == "all":
                    pg.locator('.cal .day[data-k="2026-11-02"]').click()
                    r = hscroll(pg)
                    c = clipped(pg)
                    check(not r["doc"] and not r["main"] and not r["bad"] and not c, f"{w}px сетка{(' ×' + scale) if scale else ''} {r['bad'] or ''} {c or ''}")
                    if w in (320, 390) and not scale:
                        pg.screenshot(path=str(OUT / f"w{w}-grid.png"), full_page=False)
                    sw = split_words(pg)
                    check(not sw, f"{w}px сетка и карточка дня{(' ×' + scale) if scale else ''}: без переносов внутри слов {sw or ''}")
                    pg.locator('.seg button[data-v="list"]').click()
                    one = pg.evaluate("() => [...document.querySelectorAll('.dlist b.nm')].every(b => b.getClientRects().length === 1)")
                    check(one, f"{w}px список{(' ×' + scale) if scale else ''}: фамилии без переноса")
                    pg.locator('.seg button[data-v="load"]').click()
                    r, c, sw = hscroll(pg), clipped(pg), split_words(pg)
                    one = pg.evaluate("() => [...document.querySelectorAll('.load b.nm')].every(b => b.getClientRects().length === 1)")
                    check(not r["doc"] and not r["main"] and not r["bad"] and not c and not sw and one,
                          f"{w}px нагрузка{(' ×' + scale) if scale else ''}: помещается {r['bad'] or ''} {c or ''} {sw or ''}")
                    pg.locator('.seg button[data-v="grid"]').click()
                r = hscroll(pg)
                c = clipped(pg)
                check(not r["doc"] and not r["main"] and not r["bad"] and not c,
                      f"{w}px {tab}{(' ×' + scale) if scale else ''}: без сдвига вбок и обрезки {r['bad'] or ''} {c or ''}")
                sw = split_words(pg)
                check(not sw, f"{w}px {tab}{(' ×' + scale) if scale else ''}: без переносов внутри слов {sw or ''}")
                if not scale and w in (320, 390) and tab in ("me", "all", "wish", "files"):
                    pg.screenshot(path=str(OUT / f"w{w}-{tab}.png"), full_page=False)
                if tab == "me":
                    pg.locator('[data-act="change"]').click()
                    sw = split_words(pg)
                    check(not sw and not hscroll(pg)["bad"], f"{w}px выбор себя{(' ×' + scale) if scale else ''}: без переносов внутри слов {sw or ''}")
                    pg.locator('[data-act="cancelpick"]').click()
                if tab == "me" and w in (320, 390):
                    pg.click("#helpBtn")
                    r = pg.evaluate("() => { const p = document.querySelector('#help .panel'); return p.scrollWidth > p.clientWidth; }")
                    check(not r, f"{w}px справка «?» помещается{(' ×' + scale) if scale else ''}")
                ctx.close()
    b.close()
check(not errors, f"ошибок JS нет {errors[:3]}")
srv.shutdown()
print("ПРОВАЛОВ:", len(fails))
sys.exit(1 if fails else 0)
