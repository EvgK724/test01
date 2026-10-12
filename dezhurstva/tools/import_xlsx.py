#!/usr/bin/env python3
"""График дежурств из Excel → data/schedule.json, затем сборка (index.html и календари cal/*.ics).

Запуск из корня репозитория:  python3 dezhurstva/tools/import_xlsx.py ФАЙЛ.xlsx [ЕЩЁ.xlsx …]
Формат — как у «ГРАФИК ДЕЖУРСТВ ПСО — НОЯБРЬ 2026»: строка с датами, ниже строки врачей
(«ФАМИЛИЯ телефон»), «+» — дежурство, текст в объединённых ячейках (ОТПУСК…) — отметка на эти дни,
выходные и праздники — по производственному календарю РФ (ru_calendar.py). Время смен берётся из примечания («Будни: 16:18–08:00 …»).
Месяц из файла заменяет тот же месяц в данных, остальные месяцы остаются.
Телефоны и лист пожеланий в приложение не попадают."""
import calendar, collections, datetime as dt, json, pathlib, re, sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import openpyxl
import ru_calendar

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "schedule.json"
TZ = "Asia/Yekaterinburg"
DEFAULT_SHIFT = {"work": ["16:18", "08:00"], "off": ["08:00", "08:00"]}
MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь",
          "Октябрь", "Ноябрь", "Декабрь"]
MONTH_STEMS = ["январ", "феврал", "март", "апрел", "ма", "июн", "июл", "август", "сентябр", "октябр", "ноябр", "декабр"]
TR = dict(zip("абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
              "a b v g d e e zh z i y k l m n o p r s t u f kh ts ch sh shch _ y _ e yu ya".split()))


def slug(s):
    out = "".join(TR.get(ch, ch if ch.isascii() and ch.isalnum() else "-") for ch in s.lower()).replace("_", "")
    return re.sub(r"-+", "-", out).strip("-")


def pretty(word):
    """КЛЕПИКОВ → Клепиков, ИВАНОВА-ПЕТРОВА → Иванова-Петрова; инициалы «Е.В.» не трогаем."""
    if "." in word:
        return word.upper()
    return "-".join(p[:1].upper() + p[1:].lower() for p in word.split("-"))


def split_person(text):
    """«ИВАНОВ 89000000000» → «Иванов»: телефон и прочие цифры отбрасываются."""
    name = re.sub(r"[+\d()\-–]{2,}[\d()\-– ]*", " ", str(text))
    name = re.sub(r"[\d,;:]+", " ", name)
    words = [w for w in name.split() if re.search(r"[A-Za-zА-Яа-яЁё]", w)]
    return " ".join(pretty(w) for w in words)


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def fill_of(cell):
    f = cell.fill
    if f is None or f.fill_type in (None, "none"):
        return None
    c = f.fgColor
    return c.rgb if c is not None and c.type == "rgb" else f"{c.type}:{c.theme if c.type == 'theme' else c.indexed}:{c.tint}"


def month_from_text(text):
    t = str(text).lower()
    y = re.search(r"(20\d\d)", t)
    for i, stem in enumerate(MONTH_STEMS):
        if re.search(r"\b" + stem + (r"(?:й|я|е|ь|ю)\b" if stem == "ма" else r"[а-я]*"), t) and y:
            return int(y.group(1)), i + 1
    return None


def parse_sheet(ws):
    # 1. Строка с датами — та, где больше всего дат (или чисел 1…31 при заголовке «… НОЯБРЬ 2026»)
    best = None
    title_month = None
    for r in range(1, min(ws.max_row, 15) + 1):
        for c in range(1, min(ws.max_column, 5) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and title_month is None:
                title_month = month_from_text(v)
    for r in range(1, min(ws.max_row, 15) + 1):
        cols = {}
        for c in range(2, ws.max_column + 1):
            v = ws.cell(r, c).value
            d = as_date(v)
            if d is None and title_month and isinstance(v, (int, float)) and 1 <= v <= 31 and float(v).is_integer():
                try:
                    d = dt.date(title_month[0], title_month[1], int(v))
                except ValueError:
                    d = None
            if d:
                cols[c] = d
        if len(cols) >= 28 and (best is None or len(cols) > len(best[1])):
            best = (r, cols)
    if not best:
        return None
    hr, cols = best
    months = collections.Counter((d.year, d.month) for d in cols.values())
    (year, month), _ = months.most_common(1)[0]
    cols = {c: d for c, d in cols.items() if (d.year, d.month) == (year, month)}
    if len(cols) != calendar.monthrange(year, month)[1]:
        print(f"  ! в строке дат {len(cols)} дней из {calendar.monthrange(year, month)[1]}")

    # 2. Строки врачей: ниже строки дат, до «Дежурных за день» / «Итого» / пустой строки после врачей
    doctors = []
    for r in range(hr + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        a = "" if a is None else str(a).strip()
        if re.match(r"(дежурн|итог|всего)", a, re.I):
            break
        if not a or not re.search(r"[А-Яа-яЁёA-Za-z]", a):
            if doctors:
                break
            continue
        name = split_person(a)
        if name:
            doctors.append((r, name))

    # 3. Выходные и праздники — всегда по производственному календарю РФ; заливка файла только для сверки
    cal = ru_calendar.month(year, month)
    off = {dt.date.fromisoformat(d) for d in cal["off"]}
    how = "по календарю РФ" + ("" if cal["verified"] else " (переносы этого года не сверены с постановлением!)")

    def detect(row):
        f = {c: fill_of(ws.cell(row, c)) for c in cols}
        we = collections.Counter(f[c] for c, d in cols.items() if d.weekday() >= 5).most_common(1)
        wd = collections.Counter(f[c] for c, d in cols.items() if d.weekday() < 5).most_common(1)
        if we and wd and we[0][0] != wd[0][0]:
            return {d for c, d in cols.items() if f[c] == we[0][0]}
        return None
    filled = detect(hr) or (detect(doctors[0][0]) if doctors else None)
    if filled is not None and filled != off:
        diff = sorted(filled ^ off)
        how += "; заливка в файле расходится с календарём: " + ", ".join(d.strftime("%d.%m") for d in diff)

    # 4. Примечания: время смен
    notes = " ".join(str(ws.cell(r, c).value) for r in range(1, ws.max_row + 1) for c in (1,)
                     if isinstance(ws.cell(r, c).value, str))
    shift = json.loads(json.dumps(DEFAULT_SHIFT))
    m = re.search(r"Будн[^:]*:\s*(\d{1,2}[:.]\d\d)\s*[–—-]\s*(\d{1,2}[:.]\d\d)", notes)
    if m:
        shift["work"] = [m.group(1).replace(".", ":").zfill(5), m.group(2).replace(".", ":").zfill(5)]
    m = re.search(r"Выходн[^:]*:\s*(\d{1,2}[:.]\d\d)\s*[–—-]\s*(\d{1,2}[:.]\d\d)", notes)
    if m:
        shift["off"] = [m.group(1).replace(".", ":").zfill(5), m.group(2).replace(".", ":").zfill(5)]

    # 5. Ячейки: «+» — дежурство; текст — отметка на объединённый диапазон
    merged = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = rng
    duty = collections.defaultdict(list)
    marks = []
    for r, name in doctors:
        for c, d in cols.items():
            v = ws.cell(r, c).value
            if v is None:
                continue
            s = str(v).strip()
            if not s or s.startswith("="):
                continue
            if s in ("+", "＋") or s.startswith("+"):
                duty[d.isoformat()].append(name)
                continue
            rng = merged.get((r, c))
            span = [cols[cc] for cc in range(rng.min_col, rng.max_col + 1) if cc in cols] if rng else [d]
            label = re.sub(r"[\d.\-–—/ ]+$", "", s).strip() or s
            marks.append({"name": name, "from": min(span).isoformat(), "to": max(span).isoformat(),
                          "label": label[:1].upper() + label[1:].lower()})
    return {"year": year, "month": month, "doctors": [n for _, n in doctors], "off": sorted(d.isoformat() for d in off),
            "shift": shift, "duty": dict(duty), "marks": marks, "off_how": how}


def main(paths):
    data = json.loads(DATA.read_text()) if DATA.exists() else {"tz": TZ, "doctors": {}, "months": {}}
    by_name = {v["name"]: k for k, v in data["doctors"].items()}
    by_surname = collections.defaultdict(list)
    for k, v in data["doctors"].items():
        by_surname[v["name"].split()[0]].append(k)

    def doc_id(name):
        if name in by_name:
            return by_name[name]
        surname = name.split()[0]
        if len(by_surname[surname]) == 1 and len(name.split()) == 1:
            return by_surname[surname][0]
        base = slug(surname) or "doc"
        i, did = 2, base
        while did in data["doctors"] and data["doctors"][did]["name"] != name:
            did = slug(name) if i == 2 and slug(name) not in data["doctors"] else f"{base}-{i}"
            i += 1
        data["doctors"][did] = {"name": name}
        by_name[name] = did
        by_surname[surname].append(did)
        return did

    done = 0
    for p in paths:
        wb = openpyxl.load_workbook(p)
        for ws in wb.worksheets:
            res = parse_sheet(ws)
            if not res:
                continue
            key = f"{res['year']}-{res['month']:02d}"
            order = [doc_id(n) for n in res["doctors"]]
            duty = {d: [doc_id(n) for n in names] for d, names in sorted(res["duty"].items())}
            marks = [{"who": doc_id(m.pop("name")), **m} for m in res["marks"]]
            data["months"][key] = {"title": f"{MONTHS[res['month'] - 1]} {res['year']}", "order": order,
                                   "shift": res["shift"], "off": res["off"], "duty": duty, "marks": marks}
            done += 1
            # Сводка для проверки глазами
            n_days = calendar.monthrange(res["year"], res["month"])[1]
            all_days = [dt.date(res["year"], res["month"], i).isoformat() for i in range(1, n_days + 1)]
            cnt = collections.Counter(x for v in duty.values() for x in v)
            print(f"{ws.title}: {data['months'][key]['title']}, лист «{ws.title}»")
            print("  смены: будни", "–".join(res["shift"]["work"]), "· выходные и праздники", "–".join(res["shift"]["off"]))
            hol = [d for d in res["off"] if dt.date.fromisoformat(d).weekday() < 5]
            wsat = [d for d in all_days if dt.date.fromisoformat(d).weekday() >= 5 and d not in res["off"]]
            print(f"  выходные определены {res['off_how']}; праздники в будни: {', '.join(hol) or 'нет'}"
                  + (f"; рабочие сб/вс: {', '.join(wsat)}" if wsat else ""))
            for did in order:
                print(f"  {data['doctors'][did]['name']:<16} {cnt[did]:>2}")
            empty = [d for d in all_days if not duty.get(d)]
            many = [d for d in all_days if len(duty.get(d, [])) > 1]
            if empty:
                print("  ! дни без дежурного:", ", ".join(empty))
            if many:
                print("  ! дни с несколькими дежурными:", ", ".join(many))
            for m in marks:
                print(f"  отметка: {data['doctors'][m['who']]['name']} — {m['label']} {m['from']} … {m['to']}")
    if not done:
        sys.exit("В файле не найдено строки с датами месяца — формат не похож на график дежурств.")
    data["months"] = dict(sorted(data["months"].items()))
    data["updated"] = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    DATA.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n")
    print("→", DATA.relative_to(ROOT.parent))
    import build
    build.main()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1:])
