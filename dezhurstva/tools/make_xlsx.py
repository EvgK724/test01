#!/usr/bin/env python3
"""Файл графика в оформлении образца «ГРАФИК ДЕЖУРСТВ ПСО — НОЯБРЬ 2026».

Публичный (для вкладки «Excel»):   python3 dezhurstva/tools/make_xlsx.py 2026-11
   → dezhurstva/files/2026-11.xlsx — один лист, фамилии без телефонов.
Для организатора (не в репозиторий!): python3 dezhurstva/tools/make_xlsx.py 2026-11 --full ВЫХОД.xlsx
   [--draft private/<месяц>-draft.json] [--phones phones.json] [--wishes wishes-text.json] [--note "строка" …]
   → второй лист «Пожелания и итоги»: пожелания, сутки, даты, часы, доля ставки, норма по календарю РФ.
phones.json: {"klepikov": "89…", …}; wishes.json: {"klepikov": "текст пожеланий", …}.
Выходные и праздники — по производственному календарю РФ (ru_calendar.py)."""
import argparse, calendar, datetime as dt, io, json, pathlib, re, sys, zipfile
from xml.sax.saxutils import escape

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import ru_calendar
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties

ROOT = pathlib.Path(__file__).resolve().parent.parent
GEN = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря"]
MON = ["ЯНВАРЬ", "ФЕВРАЛЬ", "МАРТ", "АПРЕЛЬ", "МАЙ", "ИЮНЬ", "ИЮЛЬ", "АВГУСТ", "СЕНТЯБРЬ", "ОКТЯБРЬ", "НОЯБРЬ", "ДЕКАБРЬ"]
INK, INK2, MUTED, TEAL, BROWN = "FF213C56", "FF203449", "FF637386", "FF126E69", "FF6F511F"
HEAD, OFF, WHITE, ZEBRA, TOTAL, VAC, NOTE = "FFD5E0EA", "FFCFD9E2", "FFFFFFFF", "FFF8FAFC", "FFE8EFF5", "FFEAD4A3", "FFFBF0D9"
thin = Side(style="thin", color="FFFFFFFF")
thin_d = Side(style="thin")
_cfg = ROOT / "data" / "config.json"
WEEK_HOURS = json.loads(_cfg.read_text()).get("week_hours", 39) if _cfg.exists() else 39
WD_RU = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]


def formula(ws, row, col, f, value):
    """Формула + её готовое значение: Excel пересчитает сам, а просмотр на iPhone покажет значение сразу."""
    c = ws.cell(row, col, f)
    ws.parent._dz_cache.setdefault(ws.title, {})[c.coordinate] = value
    return c


def fill(c):
    return PatternFill("solid", fgColor=c)


def font(sz, b=False, color=INK2):
    return Font(name="Arial", size=sz, bold=b, color=color)


C = Alignment(horizontal="center", vertical="center")
CW = Alignment(horizontal="center", vertical="center", wrap_text=True)
V = Alignment(vertical="center")
VW = Alignment(vertical="center", wrap_text=True)
# Таблица графика: тонкие чёрные границы у каждой ячейки — видны на печати (просьба организатора 26 октября 2026)
GRID = Border(left=thin_d, right=thin_d, top=thin_d, bottom=thin_d)


def hm(s):
    h, m = map(int, s.split(":"))
    return h, m


def shift_minutes(a, b):
    s, e = hm(a), hm(b)
    m = (e[0] * 60 + e[1]) - (s[0] * 60 + s[1])
    return m if m > 0 else m + 1440


def dur_text(m):
    h, mm = divmod(m, 60)
    return f"{h} ч {mm} мин" if mm else f"{h} ч"


def plural(n, one, few, many):
    a, b = n % 100, n % 10
    return many if 10 < a < 20 else few if 1 < b < 5 else one if b == 1 else many


def days_text(ds):
    return ", ".join(f"{d.day} {GEN[d.month - 1]}" for d in ds)


def build(mk, data, full=False, phones=None, wishes=None, notes=()):
    y, m = map(int, mk.split("-"))
    M = data["months"][mk]
    n = calendar.monthrange(y, m)[1]
    days = [dt.date(y, m, i) for i in range(1, n + 1)]
    cal = ru_calendar.month(y, m)
    off = set(cal["off"])
    hol = [dt.date.fromisoformat(d) for d in cal["hol"]]
    work_we = [dt.date.fromisoformat(d) for d in cal["work"]]
    short = [dt.date.fromisoformat(d) for d in cal["short"]]
    order = M["order"]
    names = {i: data["doctors"][i]["name"].upper() for i in order}
    phones = phones or {}
    lastc = n + 2
    title = f"{MON[m - 1].capitalize()} {y}"

    wb = Workbook()
    wb._dz_cache = {}
    ws = wb.active
    ws.title = title
    ws.sheet_view.showGridLines = False
    # Шапка
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=lastc)
    ws["A1"] = f"ГРАФИК ДЕЖУРСТВ ПСО — {MON[m - 1]} {y}" + (" (ПРЕДВАРИТЕЛЬНЫЙ)" if M.get("status") == "preliminary" else "")
    ws["A1"].font, ws["A1"].alignment = font(16, True, INK), C
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 30, 7
    ws.row_dimensions[3].height, ws.row_dimensions[4].height = 19, 16
    ws.merge_cells("A3:A4")
    ws["A3"] = "Врач, телефон" if full and phones else "Врач"
    ws["A3"].font, ws["A3"].alignment, ws["A3"].fill = font(11, True, INK), C, fill(HEAD)
    ws["A3"].border = Border(left=thin_d, right=thin_d, top=thin_d, bottom=thin_d)
    ws["A4"].border = Border(left=thin_d, right=thin_d, bottom=thin_d)
    for i, d in enumerate(days):
        col = i + 2
        isoff = d.isoformat() in off
        c = ws.cell(3, col, dt.datetime(d.year, d.month, d.day))
        c.number_format, c.font, c.alignment, c.fill = "d", font(11, True, INK), C, fill(OFF if isoff else HEAD)
        c.border = GRID
        c = formula(ws, 4, col, f'=CHOOSE(WEEKDAY({L(col)}3,2),"Пн","Вт","Ср","Чт","Пт","Сб","Вс")', WD_RU[d.weekday()])
        c.font, c.alignment, c.fill, c.border = font(9, True, INK), C, fill(OFF if isoff else HEAD), GRID
        ws.column_dimensions[L(col)].width = 3.7
    ws.merge_cells(start_row=3, start_column=lastc, end_row=4, end_column=lastc)
    c = ws.cell(3, lastc, "Всего")
    c.font, c.alignment, c.fill, c.border = font(11, True, INK), C, fill(HEAD), GRID
    ws.cell(4, lastc).border = GRID
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions[L(lastc)].width = 6.4

    # Врачи
    marks = {}
    for mk_ in M["marks"]:
        marks.setdefault(mk_["who"], []).append(mk_)
    first, last = 5, 5 + len(order) - 1
    for r, did in enumerate(order, start=first):
        ws.row_dimensions[r].height = 30
        base = ZEBRA if (r - first) % 2 else WHITE
        a = ws.cell(r, 1, names[did] + (f" {phones[did]}" if full and phones.get(did) else ""))
        a.font, a.alignment, a.fill = font(10.5), V, fill(base)
        a.border = GRID
        for i, d in enumerate(days):
            col = i + 2
            c = ws.cell(r, col, "+" if did in M["duty"].get(d.isoformat(), []) else None)
            c.font, c.alignment, c.border = font(17, True, TEAL), C, GRID
            c.fill = fill(OFF if d.isoformat() in off else base)
        for x in marks.get(did, []):
            a_, b_ = dt.date.fromisoformat(x["from"]), dt.date.fromisoformat(x["to"])
            c1, c2 = a_.day + 1, b_.day + 1
            text = x["label"].upper() + (f" {a_.day:02d}–{b_.day:02d}.{m:02d}" if a_ != b_ else f" {a_.day:02d}.{m:02d}")
            if c2 > c1:
                ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            c = ws.cell(r, c1, text)
            c.font, c.alignment, c.fill = font(15 if c2 - c1 >= 6 else 9, True, BROWN), CW, fill(VAC)
            for cc in range(c1, c2 + 1):
                ws.cell(r, cc).border = GRID
                ws.cell(r, cc).fill = fill(VAC)
        t = formula(ws, r, lastc, f'=COUNTIF(B{r}:{L(lastc - 1)}{r},"+")', sum(did in M["duty"].get(d.isoformat(), []) for d in days))
        t.font, t.alignment, t.fill = font(12, True, INK), C, fill(TOTAL)
        t.border = GRID
    tr = last + 1
    ws.row_dimensions[tr].height = 23
    c = ws.cell(tr, 1, "Дежурных за день")
    c.font, c.alignment, c.fill, c.border = font(10, True, INK), V, fill(HEAD), GRID
    per_day = [sum(x in order for x in M["duty"].get(d.isoformat(), [])) for d in days]
    for col in range(2, lastc):
        c = formula(ws, tr, col, f'=COUNTIF({L(col)}{first}:{L(col)}{last},"+")', per_day[col - 2])
        c.font, c.alignment, c.fill, c.border = font(10, True, INK), C, fill(HEAD), GRID
    c = formula(ws, tr, lastc, f"=SUM(B{tr}:{L(lastc - 1)}{tr})", sum(per_day))
    c.font, c.alignment, c.fill, c.border = font(10, True, INK), C, fill(HEAD), GRID
    red = dict(font=Font(color="FFA52C2C"), fill=PatternFill(bgColor="FFFCE2DF"))
    ws.conditional_formatting.add(f"B{tr}:{L(lastc - 1)}{tr}", CellIsRule(operator="notEqual", formula=["1"], **red))
    ws.row_dimensions[tr + 1].height = 7

    # Примечания под таблицей
    w_min = shift_minutes(*M["shift"]["work"])
    o_min = shift_minutes(*M["shift"]["off"])
    vac_txt = "; ".join(f"{x['label'].lower()} — {data['doctors'][x['who']]['name']} {x['from'][8:]}–{x['to'][8:]}.{m:02d}" for x in M["marks"])
    hol_txt = (", праздники " + days_text(hol)) if hol else ""
    ww_txt = (f"; рабочие по календарю РФ: {days_text(work_we)}" if work_we else "")
    lines = [f"+ — дежурство; серым выделены субботы, воскресенья{hol_txt} (производственный календарь РФ){ww_txt}"
             + (f"; {vac_txt}." if vac_txt else "."),
             f"Будни: {'–'.join(M['shift']['work'])} ({dur_text(w_min)}). Выходные и праздники: {'–'.join(M['shift']['off'])} ({dur_text(o_min)})."
             + (" Расчёт нагрузки — на втором листе." if full else "")]
    # Под таблицей — только в версии организатора; файл для оперативного отдела — одна таблица
    for i, text in enumerate(lines if full else []):
        r = tr + 2 + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=lastc)
        c = ws.cell(r, 1, text)
        c.font, c.alignment = font(9, False, MUTED), V
        ws.row_dimensions[r].height = 18
    ws.freeze_panes = "B5"
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", 9
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    for side in ("left", "right", "top", "bottom"):
        setattr(ws.page_margins, side, 0.315)

    if full:
        sheet2(wb, ws, title, mk, data, M, order, names, days, off, hol, short, cal, first, tr, lastc, w_min, wishes or {}, notes)
    wb.properties.creator = "Дежурства ПСО"
    stamp = dt.datetime.fromisoformat(data["updated"].replace("Z", "+00:00")).replace(tzinfo=None)
    wb.properties.created = wb.properties.modified = stamp
    return wb


def sheet2(wb, ws1, title, mk, data, M, order, names, days, off, hol, short, cal, first, tr, lastc, w_min, wishes, notes):
    s = wb.create_sheet("Пожелания и итоги")
    s.sheet_view.showGridLines = False
    q = f"'{title}'"
    y, m = map(int, mk.split("-"))
    s.merge_cells("A1:G1")
    s["A1"] = f"{MON[m - 1]} {y} — ПОЖЕЛАНИЯ И НАГРУЗКА"
    s["A1"].font, s["A1"].alignment = font(16, True, INK), C
    s.row_dimensions[1].height, s.row_dimensions[2].height, s.row_dimensions[3].height = 30, 7, 32
    heads = ["Врач", "Пожелания и ограничения", "Всего", "Сутки", f"Даты {GEN[m - 1]}", "Дежурства, ч:мин", "Доля ставки"]
    for i, h in enumerate(heads, 1):
        c = s.cell(3, i, h)
        c.font, c.alignment, c.fill = font(10, True, INK), CW, fill(HEAD)
        c.border = Border(left=thin if i > 1 else None, right=thin if i < 7 else None, bottom=thin)
    for k, w in zip("ABCDEFG", (20, 40, 7, 8, 28, 15, 12)):
        s.column_dimensions[k].width = w
    off_cols = [L(d.day + 1) for d in days if d.isoformat() in off]
    w_h, w_m = divmod(w_min, 60)
    norm_row = 4 + len(order) + 6  # строка «Норма, ч»: ИТОГО, пустая, примечание, пустая, часы, дни, норма
    norm = WEEK_HOURS / 5 * cal["workdays"] - len(short)
    tot = [0, 0, 0.0, 0.0]
    for j, did in enumerate(order):
        r = 4 + j
        r1 = first + j
        base = ZEBRA if j % 2 else WHITE
        s.row_dimensions[r].height = 42
        vals = [names[did], wishes.get(did, ""), f"={q}!{L(lastc)}{r1}",
                "=" + "+".join(f'COUNTIF({q}!{c}{r1},"+")' for c in off_cols) if off_cols else "=0",
                f'=IF(C{r}=0,"—",REPLACE(' + "&".join(f'IF({q}!{L(d.day + 1)}{r1}="+",", "&TEXT(DAY({q}!{L(d.day + 1)}$3),"00"),"")' for d in days) + ',1,2,""))',
                f"=(C{r}-D{r})*TIME({w_h},{w_m},0)+D{r}", f"=F{r}*24/$G${norm_row}"]
        mine = [d for d in days if did in M["duty"].get(d.isoformat(), [])]
        cnt, fu = len(mine), sum(d.isoformat() in off for d in mine)
        hours = ((cnt - fu) * w_min + fu * 1440) / 1440
        share = hours * 24 / norm if norm else 0
        cached = {3: cnt, 4: fu, 5: ", ".join(f"{d.day:02d}" for d in mine) or "—", 6: hours, 7: share}
        for k, v in zip(range(4), (cnt, fu, hours, share)):
            tot[k] += v
        for i, v in enumerate(vals, 1):
            c = formula(s, r, i, v, cached[i]) if i in cached else s.cell(r, i, v)
            c.font = font(10.5 if i in (1, 5) else 10 if i == 2 else 11)
            c.alignment = VW if i in (1, 2) else (CW if i == 5 else C)
            c.fill = fill(TOTAL if i in (6, 7) else base)
            c.border = Border(left=thin if i > 1 else None, right=thin if i < 7 else None, top=thin_d if i == 1 else thin,
                              bottom=thin_d if i == 1 else thin)
        s.cell(r, 6).number_format = "[h]:mm"
        s.cell(r, 7).number_format = "0.00"
    r = 4 + len(order)
    s.merge_cells(f"A{r}:B{r}")
    s.row_dimensions[r].height = 25
    sums = {3: tot[0], 4: tot[1], 6: tot[2], 7: tot[3]}
    for i, v in enumerate(["ИТОГО", None, f"=SUM(C4:C{r - 1})", f"=SUM(D4:D{r - 1})", f"{len(days)} {plural(len(days), 'день', 'дня', 'дней')}", f"=SUM(F4:F{r - 1})", f"=SUM(G4:G{r - 1})"], 1):
        c = formula(s, r, i, v, sums[i]) if i in sums else s.cell(r, i, v)
        c.font, c.fill, c.alignment = font(11, True, INK), fill(HEAD), V if i == 1 else C
    s.cell(r, 6).number_format, s.cell(r, 7).number_format = "[h]:mm", "0.00"
    hol_txt = ", ".join(f"{d.day} {GEN[d.month - 1]}" for d in hol)
    r += 2
    s.merge_cells(f"A{r}:E{r}")
    s.row_dimensions[r].height = 30
    c = s.cell(r, 1, f"Будни — {w_h}:{w_m:02d}; субботы, воскресенья{' и ' + hol_txt if hol_txt else ''} — 24:00. "
                     "Смена учитывается целиком по дате её начала. Выходные и праздники — по производственному календарю РФ.")
    c.font, c.alignment, c.fill = font(10, False, BROWN), VW, fill(NOTE)
    s.merge_cells(f"F{r}:G{r}")
    c = s.cell(r, 6, "НОРМА 1 СТАВКИ")
    c.font, c.alignment, c.fill = font(10, True, INK), C, fill(HEAD)
    red = dict(font=Font(color="FFA52C2C"), fill=PatternFill(bgColor="FFFCE2DF"))
    rows = [("Дней без дежурного", f"=COUNTIF({q}!B{tr}:{L(lastc - 1)}{tr},0)", "Часов в неделе", WEEK_HOURS),
            ("Дней с несколькими дежурными", f'=COUNTIF({q}!B{tr}:{L(lastc - 1)}{tr},">1")', "Рабочих дней", cal["workdays"]),
            (None, None, "Норма, ч", f"=G{r + 2}/5*G{r + 3}-G{r + 5}"),
            (None, None, "Сокращено, ч", len(short))]
    for k, (a, f_, lbl, v) in enumerate(rows):
        rr = r + 2 + k
        s.row_dimensions[rr].height = 24
        if a:
            s.merge_cells(f"A{rr}:B{rr}")
            s.cell(rr, 1, a).font = font(10)
            per = [sum(x in order for x in M["duty"].get(d.isoformat(), [])) for d in days]
            c = formula(s, rr, 3, f_, sum(n == 0 for n in per) if k == 0 else sum(n > 1 for n in per))
            c.font, c.alignment = font(11, True, TEAL), C
        bold = lbl == "Норма, ч"
        c = s.cell(rr, 6, lbl)
        c.font, c.fill, c.alignment = font(11 if bold else 10, bold, INK if bold else INK2), fill(HEAD if bold else TOTAL), V
        c = formula(s, rr, 7, v, norm) if bold else s.cell(rr, 7, v)
        c.font, c.fill, c.alignment = font(11 if bold else 10, bold, INK if bold else INK2), fill(HEAD if bold else TOTAL), C
    s.conditional_formatting.add(f"C{r + 2}:C{r + 3}", CellIsRule(operator="notEqual", formula=["0"], **red))
    info = []
    if hol:
        info.append(f"Праздники и перенесённые выходные в будни: {hol_txt}.")
    if short:
        info.append("Сокращённые предпраздничные дни: " + ", ".join(f"{d.day} {GEN[d.month - 1]}" for d in short)
                    + f"; начало дежурства — по графику, {M['shift']['work'][0]}.")
    if not cal["verified"]:
        info.append("Внимание: переносы выходных на этот год не сверены с постановлением Правительства.")
    info.append(f"Доля ставки = часы дежурств / норму. Норма: {WEEK_HOURS} / 5 × {cal['workdays']}"
                + (f" − {len(short)}" if short else "") + ". Дневная работа в расчёт нагрузки не включена.")
    info += list(notes)
    rr = r + 4
    for text in info:
        s.merge_cells(f"A{rr}:E{rr}")
        c = s.cell(rr, 1, text)
        c.font, c.alignment = font(9, False, MUTED), VW
        s.row_dimensions[rr].height = 26
        rr += 1
    s.freeze_panes = "A4"
    s.page_setup.orientation, s.page_setup.paperSize = "landscape", 9
    s.page_setup.fitToWidth, s.page_setup.fitToHeight = 1, 0
    s.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def with_values(xml, cells):
    """Вписать готовые значения формул (<v>) — openpyxl оставляет их пустыми."""
    for ref, v in cells.items():
        if isinstance(v, str):
            t, val = ' t="str"', escape(v)
        else:
            t, val = "", repr(round(float(v), 10)) if isinstance(v, float) else str(int(v))
        xml, n = re.subn(r'(<c r="%s")([^>]*)(><f>.*?</f>)<v\s*/>' % ref, lambda m: f"{m.group(1)}{m.group(2)}{t}{m.group(3)}<v>{val}</v>", xml, count=1)
        if not n:
            raise RuntimeError(f"нет формулы в ячейке {ref}")
    return xml


def save(wb, path, stamp):
    """Сохранение с готовыми значениями формул и постоянными отметками времени внутри архива."""
    buf = io.BytesIO()
    wb.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    cache = getattr(wb, "_dz_cache", {})
    sheets = {f"xl/worksheets/sheet{i + 1}.xml": cache.get(ws.title, {}) for i, ws in enumerate(wb.worksheets)}
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in zin.infolist():
            ni = zipfile.ZipInfo(info.filename, date_time=stamp.timetuple()[:6])
            ni.compress_type = zipfile.ZIP_DEFLATED
            data = zin.read(info.filename)
            if sheets.get(info.filename):
                data = with_values(data.decode(), sheets[info.filename]).encode()
            z.writestr(ni, data)
    pathlib.Path(path).parent.mkdir(parents=True, exist_ok=True)
    pathlib.Path(path).write_bytes(out.getvalue())


def public(mk, data):
    path = ROOT / "files" / f"{mk}.xlsx"
    stamp = dt.datetime.fromisoformat(data["updated"].replace("Z", "+00:00"))
    save(build(mk, data), path, stamp)
    return path


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("month")
    ap.add_argument("--data", default=str(ROOT / "data" / "schedule.json"))
    ap.add_argument("--draft", help="черновик месяца из compose.py (private/<месяц>-draft.json)")
    ap.add_argument("--full")
    ap.add_argument("--phones")
    ap.add_argument("--wishes")
    ap.add_argument("--note", action="append", default=[])
    a = ap.parse_args()
    data = json.loads(pathlib.Path(a.data).read_text())
    if a.draft:
        data["months"][a.month] = json.loads(pathlib.Path(a.draft).read_text())
        data["updated"] = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    if a.month not in data["months"]:
        sys.exit(f"нет месяца {a.month} в {a.data}")
    if a.full:
        phones = json.loads(pathlib.Path(a.phones).read_text()) if a.phones else {}
        wishes = json.loads(pathlib.Path(a.wishes).read_text()) if a.wishes else {}
        stamp = dt.datetime.fromisoformat(data["updated"].replace("Z", "+00:00"))
        save(build(a.month, data, True, phones, wishes, a.note), a.full, stamp)
        print("→", a.full)
    else:
        print("→", public(a.month, data))


if __name__ == "__main__":
    main()
